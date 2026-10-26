"""
Management Command: Ingest 39 Biological Axes × 42 Product Categories Table.
Reads '39 แกนชีวภาพ × 42 หมวดผลิตภัณฑ์ — ตารางจับคู่ต.xlsx',
populates ProductCategory & AxisProductMapping,
supports --dry-run (Gotcha 5),
enforces transaction.atomic() rollback safety,
and records CatalogImportLog.
"""

import os
import re
import json
import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone

from clinical.models_catalog import ProductCategory, AxisProductMapping, CatalogImportLog
from clinical.invariants import CANONICAL_39_AXES


class Command(BaseCommand):
    help = "Import and harmonize 39 biological axes with 42 product categories from official Excel file."

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='./39 แกนชีวภาพ × 42 หมวดผลิตภัณฑ์ — ตารางจับคู่ต.xlsx',
            help='Path to the Excel file containing product mappings'
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Perform validation and preview changes without modifying the database (Gotcha 5)'
        )

    def handle(self, *args, **options):
        file_path = options['file']
        dry_run = options['dry_run']

        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"File not found: {file_path}"))
            return

        self.stdout.write(self.style.MIGRATE_HEADING(
            f"=== Starting Harmonization Import ({'DRY-RUN' if dry_run else 'LIVE TRANSACTION'}) ==="
        ))
        self.stdout.write(f"Source file: {file_path}\n")

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            self.stderr.write(self.style.ERROR(f"Failed to open Excel workbook: {e}"))
            return

        # -------------------------------------------------------------
        # 1. Parse Sheet 'หมวดผลิตภัณฑ์ต้นฉบับ' (42 Categories)
        # -------------------------------------------------------------
        if 'หมวดผลิตภัณฑ์ต้นฉบับ' not in wb.sheetnames:
            self.stderr.write(self.style.ERROR("Sheet 'หมวดผลิตภัณฑ์ต้นฉบับ' not found in workbook."))
            return

        ws_cats = wb['หมวดผลิตภัณฑ์ต้นฉบับ']
        parsed_categories = []
        for r in range(6, ws_cats.max_row + 1):
            c1 = ws_cats.cell(r, 1).value
            if not c1 or not str(c1).strip():
                continue
            name_raw = str(c1).strip()
            desc_raw = str(ws_cats.cell(r, 2).value or '').strip()
            items_raw = str(ws_cats.cell(r, 3).value or '').strip()

            m = re.match(r'^(\d+)\.\s*(.*)', name_raw)
            cat_num = int(m.group(1)) if m else (len(parsed_categories) + 1)

            parsed_categories.append({
                'number': cat_num,
                'name': name_raw,
                'description': desc_raw,
                'items_raw': items_raw,
            })

        self.stdout.write(self.style.SUCCESS(
            f"✓ Parsed {len(parsed_categories)} product categories from 'หมวดผลิตภัณฑ์ต้นฉบับ'"
        ))

        # -------------------------------------------------------------
        # 2. Parse Sheet 'จับคู่แกนกับผลิตภัณฑ์' (57 Mapping Rows)
        # -------------------------------------------------------------
        if 'จับคู่แกนกับผลิตภัณฑ์' not in wb.sheetnames:
            self.stderr.write(self.style.ERROR("Sheet 'จับคู่แกนกับผลิตภัณฑ์' not found in workbook."))
            return

        ws_maps = wb['จับคู่แกนกับผลิตภัณฑ์']
        parsed_mappings = []
        axes_seen = set()

        for r in range(6, ws_maps.max_row + 1):
            c1 = ws_maps.cell(r, 1).value
            if not c1 or not str(c1).strip():
                continue

            axis_str = str(c1).strip()
            m_axis = re.search(r'Axis\s+(\d+)', axis_str)
            if not m_axis:
                continue
            axis_id = int(m_axis.group(1))
            axes_seen.add(axis_id)

            c2 = str(ws_maps.cell(r, 2).value or '').strip()
            m_cat = re.search(r'(\d+)\.\s*([^\n]+)', c2)
            cat_num = int(m_cat.group(1)) if m_cat else None

            # Extract scope status
            scope_status = "ตรงขอบเขต"
            if "ขยายขอบเขต" in c2:
                scope_status = "ขยายขอบเขต"
            elif "สนับสนุน" in c2:
                scope_status = "สนับสนุน"

            # Col 3: Items list and note
            c3 = str(ws_maps.cell(r, 3).value or '').strip()

            parsed_mappings.append({
                'axis_id': axis_id,
                'axis_name': axis_str,
                'category_number': cat_num,
                'category_name': c2.split('\n')[0].strip() if c2 else '',
                'scope_status': scope_status,
                'recommended_products': c3,
            })

        self.stdout.write(self.style.SUCCESS(
            f"✓ Parsed {len(parsed_mappings)} mapping rows across {len(axes_seen)} distinct biological axes"
        ))

        # Check Invariant: Must cover all 39 Canonical Axes
        missing_axes = [i for i in range(1, 40) if i not in axes_seen]
        if missing_axes:
            self.stdout.write(self.style.WARNING(f"! Note: Biological axes missing from mapping: {missing_axes}"))
        else:
            self.stdout.write(self.style.SUCCESS("✓ Invariant Verified: All 39 Biological Axes are fully covered!"))

        # -------------------------------------------------------------
        # 3. Dry-Run Execution Mode (Gotcha 5)
        # -------------------------------------------------------------
        if dry_run:
            self.stdout.write(self.style.NOTICE("\n[DRY-RUN MODE] Preview of actions to be committed to PostgreSQL:"))
            self.stdout.write(f"- Would create/update {len(parsed_categories)} ProductCategory records.")
            self.stdout.write(f"- Would create/update {len(parsed_mappings)} AxisProductMapping records.")
            self.stdout.write("\nSample Mappings:")
            for m in parsed_mappings[:5]:
                self.stdout.write(f"  Axis {m['axis_id']} -> Cat {m['category_number']} ({m['category_name']}) [{m['scope_status']}]")
            self.stdout.write(self.style.SUCCESS("\n[DRY-RUN COMPLETED] Zero database changes made."))
            return

        # -------------------------------------------------------------
        # 4. Live Execution Mode inside transaction.atomic() (Gotcha 5)
        # -------------------------------------------------------------
        cat_map = {}
        with transaction.atomic():
            # 1. Populate / Update ProductCategory
            created_cats = 0
            updated_cats = 0
            for item in parsed_categories:
                cat_obj, created = ProductCategory.objects.update_or_create(
                    category_number=item['number'],
                    defaults={
                        'name': item['name'],
                        'description': item['description'],
                        'items_raw': item['items_raw'],
                        'is_deleted': False
                    }
                )
                cat_map[item['number']] = cat_obj
                if created:
                    created_cats += 1
                else:
                    updated_cats += 1

            self.stdout.write(self.style.SUCCESS(
                f"✓ Product Categories: {created_cats} created, {updated_cats} updated."
            ))

            # 2. Populate / Update AxisProductMapping
            created_maps = 0
            updated_maps = 0
            for item in parsed_mappings:
                cat_obj = cat_map.get(item['category_number'])
                map_obj, created = AxisProductMapping.objects.update_or_create(
                    axis_id=item['axis_id'],
                    category_name=item['category_name'],
                    defaults={
                        'axis_name': item['axis_name'],
                        'category': cat_obj,
                        'scope_status': item['scope_status'],
                        'recommended_products': item['recommended_products'],
                        'is_deleted': False
                    }
                )
                if created:
                    created_maps += 1
                else:
                    updated_maps += 1

            self.stdout.write(self.style.SUCCESS(
                f"✓ Axis Product Mappings: {created_maps} created, {updated_maps} updated."
            ))

            # 3. Record CatalogImportLog
            import_log = CatalogImportLog.objects.create(
                file_name=os.path.basename(file_path),
                is_dry_run=False,
                categories_count=len(parsed_categories),
                mappings_count=len(parsed_mappings),
                summary=f"Successfully imported {len(parsed_categories)} categories and {len(parsed_mappings)} mappings across {len(axes_seen)} axes."
            )
            self.stdout.write(self.style.SUCCESS(f"✓ Recorded Audit CatalogImportLog ID: {import_log.pk}"))

        # -------------------------------------------------------------
        # 5. Harmonize with apps/clinical/data/tspi_39_axes.json
        # -------------------------------------------------------------
        json_path = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'data', 'tspi_39_axes.json')
        if os.path.exists(json_path):
            try:
                with open(json_path, 'r', encoding='utf-8') as f:
                    axes_json = json.load(f)

                # Group products by axis_id
                axis_products_dict = {}
                for m in parsed_mappings:
                    aid = m['axis_id']
                    prods = [p.strip() for p in m['recommended_products'].replace('\r', '').split('\n') if p.strip() and not p.startswith('รายชื่อตามหมวด')]
                    if aid not in axis_products_dict:
                        axis_products_dict[aid] = []
                    for prod in prods:
                        # Split by space if multiple words e.g. "KS KERRA minoza"
                        for sub_p in prod.split():
                            clean_p = sub_p.strip()
                            if len(clean_p) > 1 and clean_p not in axis_products_dict[aid]:
                                axis_products_dict[aid].append(clean_p)

                # Update modules in JSON
                for domain in axes_json.get('domains', []):
                    for ax in domain.get('axes', []):
                        aid = ax.get('axis_id')
                        if aid in axis_products_dict and axis_products_dict[aid]:
                            existing = set(ax.get('modules', []))
                            for p in axis_products_dict[aid]:
                                if p not in existing:
                                    ax.setdefault('modules', []).append(p)

                with open(json_path, 'w', encoding='utf-8') as f:
                    json.dump(axes_json, f, ensure_ascii=False, indent=4)

                self.stdout.write(self.style.SUCCESS(f"✓ Synchronized and enriched {json_path}"))
            except Exception as e:
                self.stdout.write(self.style.WARNING(f"! Could not enrich JSON: {e}"))

        self.stdout.write(self.style.SUCCESS("\n=== 39 Axes × 42 Products Harmonization Completed Successfully ==="))
