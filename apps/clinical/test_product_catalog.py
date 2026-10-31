"""
Tests for Harmonized Product Catalog & 39 Axes Mappings (Phase 4).

ใช้ไฟล์ Excel และ JSON ชั่วคราวเท่านั้น — ห้ามแตะไฟล์ต้นฉบับใน repo
(apps/clinical/data/tspi_39_axes.json และ Excel 39×42 ที่ root)
"""

import hashlib
import json
import os
import shutil
import tempfile

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from clinical.models_catalog import ProductCategory, AxisProductMapping, CatalogImportLog
from clinical.management.commands.import_39_axes_42_products import DEFAULT_AXES_JSON


def _build_fixture_workbook(path):
    """สร้าง Excel จำลองที่มีโครงสร้างเหมือนไฟล์จริง: 42 หมวด, 39 แกน (+1 แถวเสริม), ข้อมูลเริ่มแถวที่ 6"""
    wb = openpyxl.Workbook()
    ws_cats = wb.active
    ws_cats.title = 'หมวดผลิตภัณฑ์ต้นฉบับ'
    for i in range(1, 43):
        name = "Inflammation Balance" if i == 1 else f"Category {i}"
        ws_cats.cell(5 + i, 1, f"{i}. {name}")
        ws_cats.cell(5 + i, 2, f"Description {i}")
        ws_cats.cell(5 + i, 3, f"Item {i}")

    ws_maps = wb.create_sheet('จับคู่แกนกับผลิตภัณฑ์')
    row = 6
    for axis_id in range(1, 40):
        ws_maps.cell(row, 1, f"Axis {axis_id} — Test Axis")
        ws_maps.cell(row, 2, f"{axis_id}. Category {axis_id}\nตรงขอบเขต")
        ws_maps.cell(row, 3, f"FixtureProd{axis_id}")
        row += 1
    # แกน 1 จับคู่หมวดที่สองเพิ่มเติม (ขยายขอบเขต)
    ws_maps.cell(row, 1, "Axis 1 — Test Axis")
    ws_maps.cell(row, 2, "40. Category 40\nขยายขอบเขต")
    ws_maps.cell(row, 3, "FixtureExtra")
    wb.save(path)


def _sha256(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


class ProductCatalogHarmonizationTests(TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir)
        self.xlsx_path = os.path.join(self.tmpdir, 'fixture_39x42.xlsx')
        _build_fixture_workbook(self.xlsx_path)
        self.axes_json_path = os.path.join(self.tmpdir, 'tspi_39_axes.json')
        shutil.copyfile(DEFAULT_AXES_JSON, self.axes_json_path)
        self.real_json_hash = _sha256(DEFAULT_AXES_JSON)

    def _run_import(self, **kwargs):
        call_command(
            'import_39_axes_42_products',
            file=self.xlsx_path,
            axes_json=self.axes_json_path,
            stdout=open(os.devnull, 'w'),
            **kwargs,
        )

    def test_import_command_dry_run_does_not_modify_database(self):
        """Gotcha 5: คำสั่งนำเข้าด้วย --dry-run ต้องไม่บันทึกลงฐานข้อมูล"""
        initial_cats = ProductCategory.objects.count()
        initial_maps = AxisProductMapping.objects.count()

        self._run_import(dry_run=True)

        self.assertEqual(ProductCategory.objects.count(), initial_cats)
        self.assertEqual(AxisProductMapping.objects.count(), initial_maps)
        self.assertEqual(_sha256(self.axes_json_path), self.real_json_hash)

    def test_import_command_populates_categories_and_mappings_in_transaction(self):
        """Phase 4: คำสั่งนำเข้าต้องบันทึก 42 หมวดผลิตภัณฑ์ และ 39 แกนชีวภาพลง DB จริง"""
        self._run_import()

        # Verify 42 categories
        self.assertEqual(ProductCategory.objects.count(), 42)
        cat_1 = ProductCategory.objects.filter(category_number=1).first()
        self.assertIsNotNone(cat_1)
        self.assertIn("Inflammation", cat_1.name)

        # Verify 39 axes mapping
        mappings = AxisProductMapping.objects.all()
        self.assertEqual(mappings.count(), 40)
        distinct_axes = set(mappings.values_list('axis_id', flat=True))
        self.assertEqual(len(distinct_axes), 39, "All 39 biological axes must have corresponding product mappings")
        self.assertEqual(
            AxisProductMapping.objects.get(axis_id=1, category_name="40. Category 40").scope_status,
            "ขยายขอบเขต",
        )

        # Verify audit log recorded
        import_log = CatalogImportLog.objects.first()
        self.assertIsNotNone(import_log)
        self.assertFalse(import_log.is_dry_run)

    def test_import_enriches_given_json_without_touching_repo_file(self):
        """การ sync JSON ต้องเขียนเฉพาะไฟล์ที่ระบุใน --axes-json ไม่ใช่ไฟล์ต้นฉบับใน repo"""
        self._run_import()

        self.assertEqual(_sha256(DEFAULT_AXES_JSON), self.real_json_hash)
        with open(self.axes_json_path, encoding='utf-8') as f:
            axes_json = json.load(f)
        axis_1 = next(
            ax for d in axes_json['domains'] for ax in d['axes'] if ax['axis_id'] == 1
        )
        self.assertIn("FixtureProd1", axis_1['modules'])
        self.assertIn("FixtureExtra", axis_1['modules'])

    def test_missing_excel_file_raises_error(self):
        """ไฟล์ Excel ไม่พบต้อง error ชัดเจน ไม่ใช่ผ่านเงียบ ๆ"""
        with self.assertRaises(CommandError):
            call_command(
                'import_39_axes_42_products',
                file=os.path.join(self.tmpdir, 'missing.xlsx'),
                axes_json=self.axes_json_path,
            )
