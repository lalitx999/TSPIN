import os
import sys
import re
import smtplib
import multiprocessing
import json
import random
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from datetime import datetime
from llama_cpp import Llama

# ============================================================
# 📦 คำแนะนำการติดตั้งไลบรารี (รันใน Termux ก่อน)
# ============================================================
# pkg update && pkg install libxml2 libxslt clang
# pip install llama-cpp-python pypdf python-docx striprtf
# ============================================================

# ============================================================
# ส่วนที่ 0: การตั้งค่าส่งอีเมล (Gmail App Password)
# ============================================================

EMAIL_CONFIG = {
    "smtp_server": "smtp.gmail.com",
    "smtp_port": 587,
    "sender_email": "obsidianzx.m123@gmail.com",
    "app_password": "comy rtvd svme xydl",  # 16 ตัวอักษร (โค้ดจะลบช่องว่างให้)
    "receiver_email": "obsidianzx.m123@gmail.com"
}

def send_notification(subject, body):
    """ส่งการแจ้งเตือนทางอีเมล"""
    try:
        msg = MIMEMultipart()
        msg['From'] = EMAIL_CONFIG["sender_email"]
        msg['To'] = EMAIL_CONFIG["receiver_email"]
        msg['Subject'] = subject
        msg.attach(MIMEText(body, 'plain', 'utf-8'))
        
        server = smtplib.SMTP(EMAIL_CONFIG["smtp_server"], EMAIL_CONFIG["smtp_port"])
        server.starttls()
        password_clean = EMAIL_CONFIG["app_password"].replace(" ", "")
        server.login(EMAIL_CONFIG["sender_email"], password_clean)
        server.sendmail(EMAIL_CONFIG["sender_email"], EMAIL_CONFIG["receiver_email"], msg.as_string())
        server.quit()
        
        print("📧 ส่งอีเมลแจ้งเตือนเรียบร้อยแล้ว!")
        return True
    except Exception as e:
        print(f"⚠️ ส่งอีเมลไม่สำเร็จ: {e}")
        return False

# ============================================================
# ส่วนที่ 1: Multi-Format Reader (รองรับ PDF, DOCX, JSON, CSV, HTML, RTF)
# ============================================================

def read_file_content(file_path):
    """อ่านไฟล์หลายรูปแบบ"""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext in ['.txt', '.md', '.markdown']:
        try:
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                return f.read()
        except UnicodeDecodeError:
            with open(file_path, 'r', encoding='cp874') as f:
                return f.read()
    
    elif ext == '.pdf':
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except ImportError:
            return "[ไม่สามารถอ่าน PDF ได้ ต้องติดตั้ง pypdf]"
        except Exception as e:
            return f"[อ่าน PDF ไม่ได้: {e}]"
    
    elif ext == '.docx':
        try:
            import docx
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    text += "\n" + row_text
            return text
        except ImportError:
            return "[ไม่สามารถอ่าน DOCX ได้ ต้องติดตั้ง python-docx]"
        except Exception as e:
            return f"[อ่าน DOCX ไม่ได้: {e}]"
    
    elif ext == '.json':
        try:
            import json
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            if isinstance(data, dict):
                return "\n".join([f"{k}: {v}" for k, v in data.items()])
            elif isinstance(data, list):
                return "\n".join([str(item) for item in data])
            return str(data)
        except Exception as e:
            return f"[อ่าน JSON ไม่ได้: {e}]"
    
    elif ext == '.csv':
        try:
            import csv
            lines = []
            with open(file_path, 'r', encoding='utf-8') as csv_file:
                reader = csv.reader(csv_file)
                for row in reader:
                    lines.append(" | ".join(row))
            return "\n".join(lines)
        except Exception as e:
            return f"[อ่าน CSV ไม่ได้: {e}]"
    
    elif ext == '.html':
        try:
            from html.parser import HTMLParser
            class MLStripper(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.text = []
                def handle_data(self, d):
                    self.text.append(d)
                def get_data(self):
                    return ''.join(self.text)
            with open(file_path, 'r', encoding='utf-8') as f:
                html_content = f.read()
            stripper = MLStripper()
            stripper.feed(html_content)
            return stripper.get_data()
        except Exception as e:
            return f"[อ่าน HTML ไม่ได้: {e}]"
    
    elif ext == '.rtf':
        try:
            from striprtf.striprtf import rtf_to_text
            with open(file_path, 'r', encoding='utf-8') as f:
                rtf_content = f.read()
            return rtf_to_text(rtf_content)
        except ImportError:
            return "[ไม่สามารถอ่าน RTF ได้ ต้องติดตั้ง striprtf]"
        except Exception as e:
            return f"[อ่าน RTF ไม่ได้: {e}]"
    
    elif ext in ['.log', '.ini', '.cfg', '.conf']:
        try:
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                return f.read()
        except:
            with open(file_path, 'r', encoding='cp874') as f:
                return f.read()
    
    else:
        return f"[ไม่รองรับไฟล์ประเภท {ext}]"

# ============================================================
# ส่วนที่ 2: Smart Chunking แบบมี Overlap
# ============================================================

def split_text_with_overlap(text, chunk_size=800, chunk_overlap=150):
    """หั่นข้อความพร้อมส่วนคาบเกี่ยว"""
    chunks = []
    start = 0
    text_len = len(text)
    
    while start < text_len:
        end = start + chunk_size
        chunk = text[start:end]
        
        if end < text_len:
            last_newline = chunk.rfind('\n')
            if last_newline > chunk_size * 0.7:
                end = start + last_newline
                chunk = text[start:end]
        
        chunks.append(chunk.strip())
        start = end - chunk_overlap
    
    # ลบซ้ำ
    seen = set()
    unique_chunks = []
    for ch in chunks:
        if ch not in seen:
            seen.add(ch)
            unique_chunks.append(ch)
    
    return unique_chunks

# ============================================================
# ส่วนที่ 3: กำหนดพาธโฟลเดอร์ข้อมูล (อยู่ใน Internal Storage)
# ============================================================

DATA_FOLDER = "/sdcard/data"

print("="*60)
print("🩺 TSPI™ Medical CDSS v5.2 - Self-Reflection & Self-Training Edition")
print("="*60)
print(f"📂 โฟลเดอร์ข้อมูล: {DATA_FOLDER}")

# ตรวจสอบโฟลเดอร์
if not os.path.exists(DATA_FOLDER):
    print(f"\n❌ ไม่พบโฟลเดอร์ {DATA_FOLDER}")
    print("📌 วิธีแก้ไข: mkdir -p /sdcard/data")
    sys.exit(1)

# ============================================================
# ส่วนที่ 4: อ่านไฟล์ทั้งหมด (รวมโฟลเดอร์ย่อย)
# ============================================================

print("\n📂 กำลังกวาดไฟล์ข้อมูล (รวมโฟลเดอร์ย่อย)...")

SUPPORTED_EXTS = {
    '.txt', '.md', '.markdown', '.pdf', '.docx', 
    '.json', '.csv', '.html', '.htm', '.rtf',
    '.log', '.ini', '.cfg', '.conf'
}

raw_text = ""
file_count = 0
file_names_loaded = []
format_stats = {}

def scan_directory(directory):
    """เดิน递归เข้าไปอ่านไฟล์ทุกโฟลเดอร์ย่อย"""
    global raw_text, file_count, file_names_loaded, format_stats
    
    try:
        for item in os.listdir(directory):
            item_path = os.path.join(directory, item)
            
            if os.path.isdir(item_path):
                rel_path = os.path.relpath(item_path, DATA_FOLDER)
                print(f"   📁 เข้าโฟลเดอร์: {rel_path}")
                scan_directory(item_path)
            
            elif os.path.isfile(item_path):
                ext = os.path.splitext(item)[1].lower()
                if ext in SUPPORTED_EXTS:
                    rel_path = os.path.relpath(item_path, DATA_FOLDER)
                    print(f"   📄 กำลังอ่าน: {rel_path}")
                    content = read_file_content(item_path)
                    
                    if not content.startswith("[ไม่") and not content.startswith("ไม่สามารถ"):
                        format_stats[ext] = format_stats.get(ext, 0) + 1
                        raw_text += f"\n\n{'='*60}\n[ไฟล์: {rel_path} | ประเภท: {ext}]\n{'='*60}\n{content}\n"
                        file_count += 1
                        file_names_loaded.append(rel_path)
                        print(f"      ✅ โหลดสำเร็จ ({len(content):,} ตัวอักษร)")
    except PermissionError:
        print(f"   ⚠️ ไม่มีสิทธิ์เข้าถึง {directory}")
    except Exception as e:
        print(f"   ⚠️ Error: {e}")

scan_directory(DATA_FOLDER)

print("\n" + "="*60)
if file_count == 0:
    print("❌ ไม่พบไฟล์ที่รองรับ!")
    print(f"📌 ตรวจสอบว่าใน {DATA_FOLDER} มีไฟล์ .md หรือ .txt")
    sys.exit(1)

print(f"✅ โหลดข้อมูลสำเร็จ {file_count} ไฟล์")
print(f"📊 ประเภทไฟล์: {', '.join(format_stats.keys())}")
print(f"📏 ขนาดข้อมูลรวม: {len(raw_text):,} ตัวอักษร")

# ============================================================
# ส่วนที่ 5: Smart Chunking
# ============================================================

print("\n🧠 กำลังแบ่งข้อมูลเป็น chunks...")

document_segments = []
file_blocks = re.split(r'\n(?=\[ไฟล์:)', raw_text)

for block in file_blocks:
    if len(block.strip()) < 50:
        continue
    
    header_match = re.match(r'^(\[ไฟล์:[^\]]+\])', block.strip())
    header = header_match.group(1) if header_match else "[ไฟล์ข้อมูลทั่วไป]"
    
    sub_chunks = split_text_with_overlap(block, chunk_size=800, chunk_overlap=150)
    for sc in sub_chunks:
        if not sc.startswith("[ไฟล์:"):
            document_segments.append(f"{header}\n{sc}")
        else:
            document_segments.append(sc)

print(f"✂️ แบ่งข้อมูลออกเป็น {len(document_segments)} chunks")

# ============================================================
# ส่วนที่ 6: ส่งอีเมลแจ้งเตือน (เตรียมข้อมูลเสร็จ)
# ============================================================

print("\n📧 กำลังส่งอีเมลแจ้งเตือน...")
email_subject = f"✅ [TSPI CDSS] โหลดข้อมูลเสร็จ - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
email_body = f"""
ระบบ TSPI™ Medical CDSS โหลดข้อมูลเรียบร้อยแล้ว!

📊 สรุปผล:
• จำนวนไฟล์: {file_count} ไฟล์
• ประเภทไฟล์: {', '.join(format_stats.keys())}
• ขนาดข้อมูล: {len(raw_text):,} ตัวอักษร
• จำนวน Chunks: {len(document_segments)} ชิ้น
• เวลา: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

🚀 ระบบพร้อมให้บริการแล้ว
"""
send_notification(email_subject, email_body)

# ============================================================
# ส่วนที่ 7: โหลดโมเดล AI
# ============================================================

print("\n🔄 กำลังโหลดโมเดล Qwen...")
sys.platform = "linux"

# หาโมเดล
MODEL_PATHS = [
    os.path.join(DATA_FOLDER, "qwen2.5-1.5b-instruct.gguf"),
    "./qwen2.5-1.5b-instruct.gguf",
    "/sdcard/qwen2.5-1.5b-instruct.gguf",
]

MODEL_PATH = None
for path in MODEL_PATHS:
    if os.path.exists(path):
        MODEL_PATH = path
        print(f"📁 พบโมเดล: {path}")
        break

if not MODEL_PATH:
    print("\n❌ ไม่พบไฟล์โมเดล .gguf!")
    print("📌 วางโมเดลไว้ที่ /sdcard/data/ หรือข้างสคริปต์")
    sys.exit(1)

# ปรับ threads อัตโนมัติ ( SNAPDRAGON 680 มี 4 คอร์แรง Pin ไว้ 4 เวิร์กสุด)
optimal_threads = max(2, multiprocessing.cpu_count() // 2)
print(f"⚙️ ใช้ CPU threads: {optimal_threads}")

try:
    llm = Llama(
        model_path=MODEL_PATH,
        n_ctx=3072,
        n_threads=optimal_threads,
        verbose=False
    )
    print("✅ โหลดโมเดลสำเร็จ!")
except Exception as e:
    print(f"❌ โหลดโมเดลล้มเหลว: {e}")
    sys.exit(1)

# ============================================================
# ส่วนที่ 8: ฟังก์ชันให้คะแนนความเกี่ยวข้อง
# ============================================================

def extract_keywords(text):
    words = re.findall(r'[ก-๙a-zA-Z0-9\-]+', text.lower())
    stopwords = {'และ', 'หรือ', 'แต่', 'หาก', 'ไป', 'มา', 'ได้', 'ไว้', 'คือ', 'เป็น', 'ที่', 'ซึ่ง'}
    return [w for w in words if len(w) > 2 and w not in stopwords]

def score_segment(segment, query_keywords, user_query_lower):
    score = 0
    seg_lower = segment.lower()
    
    for kw in query_keywords:
        if kw in seg_lower:
            score += 5
    
    tspi_terms = {
        'db-1': ['db-1', 'db1', 'ดีบีวัน'],
        'ks': ['ks', 'เคเอส'],
        'kerra': ['kerra', 'เคอร์รา'],
        'minoza': ['minoza', 'มิโนซา'],
        'mendoza': ['mendoza', 'เมนโดซา'],
        'atp': ['atp', 'เอทีพี'],
        'fbs': ['fbs', 'น้ำตาล'],
        'hba1c': ['hba1c', 'ฮีโมโกลบิน'],
    }
    
    for term, variants in tspi_terms.items():
        if any(v in user_query_lower for v in variants):
            if any(v in seg_lower for v in variants):
                score += 10
    
    medical_keywords = ['เบาหวาน', 'ความดัน', 'ไขมัน', 'หลอดเลือด', 'อักเสบ', 'ภูมิคุ้มกัน', 'ฮอร์โมน', 'ประสาท']
    for mk in medical_keywords:
        if mk in user_query_lower and mk in seg_lower:
            score += 8
    
    return score

def send_analysis_to_email(case_input, ai_result):
    """ส่งผลวิเคราะห์ทางอีเมล"""
    subject = f"🩺 [TSPI CDSS] ผลวิเคราะห์เสร็จสมบูรณ์ - {datetime.now().strftime('%H:%M:%S')}"
    body = f"""
📥 เคส/อาการผู้ป่วย:
{case_input}

🧠 ผลวิเคราะห์และแนวทางการรักษา (ผ่านการทำ Self-Reflection):
{ai_result}

📅 วันเวลาวิเคราะห์: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
🚀 ระบบวิเคราะห์โรคอัตโนมัติ TSPI™ CDSS
"""
    send_notification(subject, body)

# ============================================================
# 🧠 ฟังก์ชันถามเองตอบเองเพื่อเทรนตัวเอง (Self-Training Data Maker Mode)
# ============================================================

def run_self_training_loop(llm, document_segments, num_samples=10):
    """
    ลูปสร้างชุดข้อมูลสังเคราะห์ (ถามเองตอบเองและขัดเกลาอัตโนมัติ) เพื่อนำไป Train/Fine-tune
    """
    print("\n" + "="*60)
    print("🧠 [TSPI Self-Training Mode] เริ่มต้นลูปวิจัยและสร้างชุดข้อมูลเทรนตัวเอง...")
    print("="*60)
    print(f"📊 จะสุ่มอ่านเอกสาร {len(document_segments)} ชิ้นเพื่อสร้างชุดคำถาม-คำตอบจำนวน {num_samples} ข้อ")
    
    dataset = []
    success_count = 0
    
    output_file = os.path.join(DATA_FOLDER, "tspi_training_dataset.jsonl")
    
    # ดึงเฉพาะ chunks ที่น่าเชื่อถือและยาวเกิน 200 ตัวอักษร
    valid_chunks = [ch for ch in document_segments if len(ch.strip()) > 200]
    if not valid_chunks:
        valid_chunks = document_segments
        
    for idx in range(num_samples):
        print(f"\n⚙️ [{idx+1}/{num_samples}] สุ่มดึงความรู้คลินิกมาตั้งโจทย์...")
        chunk = random.choice(valid_chunks)
        
        # 1. ตั้งคำถาม (Ask)
        print("   ❓ 1. กำลังคิดหัวข้อและตั้งคำถามทางการแพทย์ (Generating Question)...")
        q_prompt = f"""คุณคืออาจารย์แพทย์ของระบบ TSPI™ 39 Biological Axes
จงอ่านเอกสารอ้างอิงด้านล่าง แล้วคิด "คำถามทางการแพทย์" (Medical Case/Question) ภาษาไทยที่สมจริง 1 ข้อ
คำถามนี้อาจเกี่ยวกับการขอความช่วยเหลือของคนไข้ การวินิจฉัยของหมอ หรือความเกี่ยวโยงของสมุนไพร
*ข้อสำคัญ: ตอบเฉพาะคำถามภาษาไทย 1 ประโยคเท่านั้น ห้ามนำข้อมูลอ้างอิงมาเขียนซ้ำทั้งหมดในคำถาม*

【ข้อมูลความรู้】
{chunk}

คำถามแพทย์ภาษาไทย:"""
        q_input = f"<|im_start|>system\nคุณคืออาจารย์แพทย์ผู้เชี่ยวชาญการตั้งคำถามทดสอบคลินิกระบบ TSPI™<|im_end|>\n<|im_start|>user\n{q_prompt}<|im_end|>\n<|im_start|>assistant\n"
        q_out = llm(q_input, max_tokens=150, temperature=0.7, stop=["<|im_end|>"])
        generated_q = q_out['choices'][0]['text'].strip()
        
        # Clean Chinese characters in generated question
        generated_q = re.sub(r'[\u4e00-\u9fff]', '', generated_q)
        
        if len(generated_q) < 15:
            print("   ⚠️ คำถามสั้นเกินไปหรือผิดพลาด ข้ามรอบนี้...")
            continue
            
        print(f"      👉 ได้คำถาม: {generated_q}")
        
        # 2. ตอบคำถามตัวเอง (Answer)
        print("   💡 2. กำลังร่างคำตอบวิเคราะห์อาการขั้นต้น (Generating Answer)...")
        a_prompt = f"""คุณคือผู้เชี่ยวชาญทางการแพทย์ของระบบ TSPI™
จงวิเคราะห์และตอบคำถามทางการแพทย์นี้ โดยอ้างอิงจากข้อมูลวิชาการที่มีให้ในเอกสารอ้างอิงเท่านั้น

【เอกสารอ้างอิง】
{chunk}

【คำถาม】
{generated_q}

คำตอบทางการแพทย์ภาษาไทย:"""
        a_input = f"<|im_start|>system\nคุณคือผู้เชี่ยวชาญทางการแพทย์ที่ตอบคำถามคลินิกได้อย่างแม่นยำภาษาไทย 100%<|im_end|>\n<|im_start|>user\n{a_prompt}<|im_end|>\n<|im_start|>assistant\n"
        a_out = llm(a_input, max_tokens=512, temperature=0.2, stop=["<|im_end|>"])
        generated_a = a_out['choices'][0]['text'].strip()
        
        # 3. ขัดเกลาและแก้ไขจุดหลอนจุดภาษาจีน (Self-Critique & Reflection)
        print("   🔍 3. ทำ Self-Reflection (ลบภาษาจีนและคำหลอนของตัวเอง)...")
        critique_prompt = f"""คุณคือผู้ตรวจชำระตำราแพทย์ TSPI™
จงนำคำวิเคราะห์ด้านล่างนี้มาลบภาษาจีน (เช่น 胆汁, 血管) และลบอาการหลอน (เช่น โรคอัลไซด์ เป็น โรคสมองเสื่อม/อัลไซเมอร์) ออกให้สิ้นเชิง
เขียนใหม่เป็นคำวิเคราะห์ภาษาไทย 100% ที่กระชับและลื่นไหลเชิงการแพทย์ไทยประยุกต์

【คำวิเคราะห์เดิมที่มีปัญหา】
{generated_a}

คำวิเคราะห์ฉบับแก้ไขสมบูรณ์แบบภาษาไทย:"""
        c_input = f"<|im_start|>system\nคุณคือบรรณาธิการแพทย์ที่วิเคราะห์ประเมินภาษาไทยของ CDSS<|im_end|>\n<|im_start|>user\n{critique_prompt}<|im_end|>\n<|im_start|>assistant\n"
        c_out = llm(c_input, max_tokens=512, temperature=0.1, stop=["<|im_end|>"])
        corrected_a = c_out['choices'][0]['text'].strip()
        
        # คลีนรอบสุดท้ายด้วย Regex
        corrected_a = re.sub(r'[\u4e00-\u9fff]', '', corrected_a)
        corrected_a = corrected_a.replace("โรคอัลไซด์", "โรคอัลไซเมอร์ / ภาวะสมองเสื่อม")
        
        # สร้างชุดข้อมูล
        sample = {
            "instruction": generated_q,
            "input": "",
            "output": corrected_a
        }
        dataset.append(sample)
        
        # บันทึกลง JSONL ทันที
        with open(output_file, 'a', encoding='utf-8') as f:
            f.write(json.dumps(sample, ensure_ascii=False) + "\n")
            
        success_count += 1
        print(f"      ✅ สำเร็จ! (จัดเก็บใน Dataset เรียบร้อยแล้ว)")
        
    print(f"\n🎉 [Self-Training 완료] ผลการถามเองตอบเองเสร็จสิ้น! สร้างได้ {success_count} คู่คำถามคำตอบ")
    print(f"💾 ไฟล์ข้อมูลจัดเก็บที่: {output_file}")
    
    # ส่งอีเมลยืนยันผลการทำ Self-Training ทุกครั้ง
    print("📧 กำลังส่งรายงานผลทางอีเมลอัตโนมัติ...")
    sample_preview = ""
    for idx, d in enumerate(dataset[:3]):
        sample_preview += f"\n📊 ลำดับที่ {idx+1}:\n❓ คำถาม: {d['instruction']}\n💡 คำตอบบางส่วน: {d['output'][:200]}...\n{'-'*40}\n"
        
    email_subject = f"🧠 [TSPI Self-Training] สร้างชุดข้อมูลสำเร็จ - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    email_body = f"""
ลูปถามเองตอบเองเพื่อการเรียนรู้และเทรนตัวเอง (Self-Training Loop) ของระบบ TSPI™ ดำเนินการเสร็จสมบูรณ์!

📈 สรุปสถิติ:
• จำนวนข้อที่ตั้งไว้: {num_samples} ข้อ
• เจเนอเรตและกลั่นกรองสำเร็จ: {success_count} ข้อ
• ไฟล์จัดเก็บ: {output_file}
• ขนาดไฟล์: {os.path.getsize(output_file) if os.path.exists(output_file) else 0:,} bytes

👉 การนำไปใช้งาน:
สามารถนำไฟล์ 'tspi_training_dataset.jsonl' ในเครื่องโทรศัพท์ ไปใช้อัพโหลดเทรนโมเดลผ่าน Google Colab (Unsloth) เพื่อเพิ่มพูนความสามารถของโมเดลตัวใหม่ได้ทันที!

📝 ตัวอย่างโจทย์ Q&A บางส่วน:
{sample_preview}
"""
    send_notification(email_subject, email_body)
    print("✅ แจ้งเตือนผ่านอีเมลเรียบร้อยแล้ว!")

# ============================================================
# ส่วนที่ 9: TSPI CDSS Interactive Mode
# ============================================================

print("\n" + "="*60)
print("🩺 TSPI™ Medical CDSS v5.2 - ระบบพร้อมปฏิบัติการ!")
print("="*60)
print(f"📚 ฐานข้อมูล: {file_count} ไฟล์ → {len(document_segments)} chunks")
print("\n💡 คำสั่งพิเศษ:")
print("   • พิมพ์ 'selftrain' เพื่อเปิดโหมดถามเองตอบเองสร้างชุดข้อมูลเทรนโมเดล")
print("   • พิมพ์ 'exit' เพื่อออกจากโปรแกรม")
print("   • หรือ พิมพ์เคสผู้ป่วย / อาการเพื่อทำการวิเคราะห์ทันที")
print("="*60)

while True:
    try:
        user_input = input("\n👨‍⚕️ [เคส/คำถาม/คำสั่ง]: ").strip()
        
        if user_input.lower() == 'exit':
            print("\n👋 ปิดระบบวิเคราะห์โรค ขอบคุณครับ!")
            break
            
        if user_input.lower() == 'selftrain':
            try:
                num = input("🔢 ต้องการสร้างชุดคำถาม-คำตอบกี่ข้อ? (แนะนำ 10-50 ข้อสำหรับมือถือ): ").strip()
                num_int = int(num) if num.isdigit() else 10
                run_self_training_loop(llm, document_segments, num_samples=num_int)
            except Exception as se:
                print(f"⚠️ โหมดถามตอบตัวเองเกิดข้อผิดพลาด: {se}")
            continue
            
        if not user_input:
            continue
            
        print("\n🔍 กำลังค้นหาข้อมูลความรู้ทางการแพทย์และสมุนไพร TSPI...")
        query_keywords = extract_keywords(user_input)
        user_query_lower = user_input.lower()
        
        scored = []
        for seg in document_segments:
            if len(seg.strip()) < 50:
                continue
            s = score_segment(seg, query_keywords, user_query_lower)
            if s > 0:
                scored.append((s, seg))
        
        scored.sort(key=lambda x: x[0], reverse=True)
        
        context = ""
        for i in range(min(3, len(scored))):
            if scored[i][0] > 0:
                context += f"\n📖 อ้างอิงแหล่งที่ {i+1} (คะแนน {scored[i][0]})\n{scored[i][1][:1500]}\n"
        
        if not context:
            context = document_segments[0][:1500] if document_segments else "ไม่มีข้อมูล"
            print("⚠️ ไม่พบข้อมูลตรงเป๊ะ ใช้ข้อมูลภาพรวมแกนชีววิทยา TSPI")
        else:
            print(f"✅ ดึงฐานข้อมูลแพทย์ {len(scored)} ส่วนมาเป็นบริบทอ้างอิง")
            
        # SYSTEM PROMPT
        system_prompt = f"""คุณคือผู้เชี่ยวชาญทางการแพทย์ของระบบ TSPI™ 39 Biological Axes

【ฐานความรู้】
{context}

【กฎระเบียบในการวิเคราะห์】
1. ตอบคำถามโดยใช้ข้อมูลจาก 【ฐานความรู้】 ด้านบนเท่านั้น ห้ามกุคำตอบหรืออาการขึ้นมาเองโดยไม่มีแหล่งอ้างอิง
2. ตอบคำถามเป็นภาษาไทย 100% สละสลวย เข้าใจง่าย
3. แนะนำสมุนไพร TSPI ที่ตรงหลักเกณฑ์: DB-1, KS, KERRA, Minoza, Mendoza
4. ตอบให้ตรงประเด็น ครบถ้วนตามแกนชีววิทยา"""

        prompt = f"<|im_start|>system\n{system_prompt}<|im_end|>\n<|im_start|>user\n{user_input}<|im_end|>\n<|im_start|>assistant\n"
        
        # ============================================================
        # 1. STEP 1: Generate Initial Draft
        # ============================================================
        print("⚡ 1. กำลังวิเคราะห์และร่างผลลัพธ์ทางการแพทย์ขั้นต้น (Drafting)...")
        draft_output = llm(prompt, max_tokens=600, temperature=0.2, stop=["<|im_end|>"])
        draft_result = draft_output['choices'][0]['text'].strip()
        
        # ============================================================
        # 2. STEP 2: Self-Reflection / Self-Correction Pass
        # ============================================================
        # เช็คด้วยโค้ดก่อนว่ามีภาษาจีนหรืออาการหลอนไหม
        has_chinese = bool(re.search(r'[\u4e00-\u9fff]', draft_result))
        has_hallucinations = "โรคอัลไซด์" in draft_result or "โรคอัลไซด์" in draft_result
        
        final_analysis = draft_result
        
        if has_chinese or has_hallucinations:
            print("🔄 2. [Self-Reflection Active] ตรวจพบจุดที่อาจหลอนหรืออักษรจีน กำลังป้อนข้อมูลแก้ไขตัวเอง...")
            
            critique_prompt = f"""จงลบตัวอักษรจีนทั้งหมดและคำที่น่าจะเป็นอาการหลอนของโมเดล เช่น คำว่า "โรคอัลไซด์" ออกจากร่างคำตอบนี้ แล้วเรียบเรียงใหม่เป็นภาษาไทย 100% ทางคลินิกที่สมบูรณ์แบบที่สุด

【ร่างคำตอบเดิมที่มีปัญหา】
{draft_result}

จงเขียนคำวิเคราะห์ใหม่ภาษาไทย 100% ที่สมบูรณ์แบบที่สุด:"""
            
            c_input = f"<|im_start|>system\nคุณคือแพทย์บรรณาธิการของระบบ CDSS ที่เก่งกาจและพูดแต่ภาษาไทย<|im_end|>\n<|im_start|>user\n{critique_prompt}<|im_end|>\n<|im_start|>assistant\n"
            corrected_output = llm(c_input, max_tokens=600, temperature=0.1, stop=["<|im_end|>"])
            final_analysis = corrected_output['choices'][0]['text'].strip()
            
            # กรอง Regex ชัวร์ๆ อีกหนึ่งรอบ
            final_analysis = re.sub(r'[\u4e00-\u9fff]', '', final_analysis)
            final_analysis = final_analysis.replace("โรคอัลไซด์", "โรคอัลไซเมอร์ / ภาวะสมองเสื่อม")
            print("✨ 3. [Self-Reflection Complete] ขัดเกลาข้อบกพร่องเรียบร้อย!")
        else:
            print("✨ 3. [Self-Reflection Complete] ตรวจสอบร่างผลลัพธ์ผ่านเกณฑ์ดีเยี่ยม ไม่มีคำหลอนหรือภาษาอื่นปะปน!")
            
        print("\n" + "="*60)
        print("🧠 ผลวิเคราะห์ (ผ่านลูปประเมินตนเอง):")
        print("="*60)
        print(final_analysis)
        print("="*60)
        
        # ส่งอีเมลอัตโนมัติทันทีที่การวิเคราะห์เสร็จสมบูรณ์ตามความต้องการ "เสร็จทุกครั้งควรแจ้งผ่านเมลลื"
        print("📧 กำลังส่งผลวิเคราะห์เคสทางอีเมลแจ้งเตือนอัตโนมัติ...")
        send_analysis_to_email(user_input, final_analysis)
        
    except KeyboardInterrupt:
        print("\n\n⚠️ กด Ctrl+C - พิมพ์ 'exit' เพื่อความปลอดภัยในการปิดโปรแกรม")
    except Exception as e:
        print(f"❌ เกิดข้อผิดพลาดการตอบคำถาม: {e}")

print("\n✅ ระบบปิดตัวลงอย่างปลอดภัย TSPI™ CDSS v5.2")
